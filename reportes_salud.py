# -*- coding: utf-8 -*-
"""Reportes en Excel del módulo Salud de estudiantes (alertas, nutrición y derivaciones)."""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "0F3460"
GRIS = "F3F6F9"
LINEA = Side(style="thin", color="DDE3EA")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)

ESTADOS = {
    "vencida": ("Vencida", "FEE2E2", "B91C1C"),
    "urgente": ("Próximos 30 días", "FFEDD5", "C2410C"),
    "proxima": ("En 31 a 90 días", "FEF3C7", "B45309"),
    "este_anio": ("Dentro de 12 meses", "E0ECFF", "1D4ED8"),
    "en_plazo": ("Al día", "DCFCE7", "15803D"),
    "fecha_a_revisar": ("Fecha a revisar", "EDE9FE", "6D28D9"),
    "sin_fecha": ("Sin fecha", "EEF2F6", "475569"),
}
CLASES = ["Normal", "Sobrepeso", "Obesidad", "Bajo peso"]


def _f(iso):
    if not iso:
        return None
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _rut(r):
    n = "".join(ch for ch in str(r or "").upper() if ch.isdigit() or ch == "K")
    if len(n) < 2:
        return r or ""
    c = f"{int(n[:-1]):,}".replace(",", ".") if n[:-1].isdigit() else n[:-1]
    return f"{c}-{n[-1]}"


def _titulo(ws, texto, sub, ancho):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
    ws.cell(1, 1, texto).font = Font(bold=True, size=15, color=AZUL)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
    ws.cell(2, 1, sub).font = Font(size=10, color="5B6B7A")
    ws.row_dimensions[1].height = 24


def _tabla(ws, fila, encabezados, filas, anchos=None):
    for j, h in enumerate(encabezados, 1):
        c = ws.cell(fila, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDE
    ws.row_dimensions[fila].height = 30
    for i, r in enumerate(filas, fila + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            c.border = BORDE
            c.alignment = Alignment(vertical="center", wrap_text=isinstance(v, str) and len(v) > 40)
            if (i - fila) % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GRIS)
            if isinstance(v, date):
                c.number_format = "DD/MM/YYYY"
    if anchos:
        for j, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.auto_filter.ref = f"A{fila}:{get_column_letter(len(encabezados))}{fila + max(1, len(filas))}"
    ws.freeze_panes = ws.cell(fila + 1, 1)
    return fila + len(filas) + 1


def _hoja_resumen(wb, datos, filtros):
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, "Reporte de salud de estudiantes · Cardiohome",
            f"Generado el {date.today().strftime('%d/%m/%Y')}", 4)
    fila = 4
    for k, v in filtros.items():
        ws.cell(fila, 1, k).font = Font(bold=True, color="5B6B7A")
        ws.cell(fila, 2, v)
        fila += 1
    fila += 1
    if "reev" in datos:
        r = datos["reev"]
        ws.cell(fila, 1, "Alertas de reevaluación").font = Font(bold=True, size=12, color=AZUL)
        fila += 1
        for k in ("vencida", "urgente", "proxima", "este_anio", "en_plazo", "fecha_a_revisar", "sin_fecha"):
            lab, bg, fg = ESTADOS[k]
            a = ws.cell(fila, 1, lab)
            a.fill = PatternFill("solid", fgColor=bg)
            a.font = Font(bold=True, color=fg)
            ws.cell(fila, 2, r["conteo"].get(k, 0))
            fila += 1
        ws.cell(fila, 1, "Total estudiantes (por especialidad)").font = Font(bold=True)
        ws.cell(fila, 2, r["total_estudiantes"])
        fila += 2
    if "nutri" in datos:
        n = datos["nutri"]
        ws.cell(fila, 1, f"Estado nutricional {n['anio']}").font = Font(bold=True, size=12, color=AZUL)
        fila += 1
        tot = n["clasificados"] or 1
        for k in CLASES:
            ws.cell(fila, 1, k)
            ws.cell(fila, 2, n["distribucion"][k])
            c = ws.cell(fila, 3, n["distribucion"][k] / tot)
            c.number_format = "0.0%"
            fila += 1
        ws.cell(fila, 1, "Estudiantes medidos").font = Font(bold=True)
        ws.cell(fila, 2, n["clasificados"])
        fila += 2
    if "deriv" in datos:
        d = datos["deriv"]
        ws.cell(fila, 1, "Derivaciones por especialidad").font = Font(bold=True, size=12, color=AZUL)
        fila += 1
        for e in d["especialidades"]:
            ws.cell(fila, 1, e["especialidad"])
            ws.cell(fila, 2, e["estudiantes"])
            c = ws.cell(fila, 3, (e["pct"] or 0) / 100)
            c.number_format = "0.0%"
            fila += 1
        ws.cell(fila, 1, "Estudiantes con al menos una derivación").font = Font(bold=True)
        ws.cell(fila, 2, d["derivados"])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12


def _hoja_reev(wb, r):
    ws = wb.create_sheet("Reevaluaciones")
    _titulo(ws, "Reevaluaciones por estudiante",
            f"Días calculados al {date.fromisoformat(r['hoy']).strftime('%d/%m/%Y')}. Una fila por estudiante y especialidad (la evaluación más reciente).", 8)
    filas = []
    for x in r["lista"]:
        lab = ESTADOS[x["estado"]][0]
        filas.append([x["nombre"], _rut(x["rut"]), x["colegio"], x["especialidad"], _f(x["fecha_evaluacion"]),
                      _f(x["fecha_reevaluacion"]), lab, x["dias"]])
    fin = _tabla(ws, 4, ["Estudiante", "RUT", "Colegio", "Especialidad", "Evaluado", "Reevaluar antes de", "Estado", "Días que faltan"],
                 filas, [32, 14, 34, 18, 12, 16, 20, 14])
    for i in range(5, fin):
        est = ws.cell(i, 7).value
        for k, (lab, bg, fg) in ESTADOS.items():
            if lab == est:
                ws.cell(i, 7).fill = PatternFill("solid", fgColor=bg)
                ws.cell(i, 7).font = Font(bold=True, color=fg)
                break
    ws2 = wb.create_sheet("Reevaluaciones por mes")
    _titulo(ws2, "Reevaluaciones por mes", "Cantidad de estudiantes que deben reevaluarse en cada mes", 2)
    _tabla(ws2, 4, ["Mes", "Estudiantes"], [[m["mes"], m["n"]] for m in r["por_mes"]], [12, 14])
    if r["por_mes"]:
        ch = BarChart()
        ch.title, ch.height, ch.width, ch.legend = "Reevaluaciones por mes", 8, 22, None
        ch.add_data(Reference(ws2, min_col=2, min_row=4, max_row=4 + len(r["por_mes"])), titles_from_data=True)
        ch.set_categories(Reference(ws2, min_col=1, min_row=5, max_row=4 + len(r["por_mes"])))
        ws2.add_chart(ch, "D4")


def _hoja_nutri(wb, n):
    ws = wb.create_sheet("Nutrición por colegio")
    _titulo(ws, f"Estado nutricional por colegio · {n['anio']}", "Una medición por estudiante. Clasificación registrada por la doctora.", 8)
    filas = [[c["colegio"], c["total"], c["Normal"], c["Sobrepeso"], c["Obesidad"], c["Bajo peso"],
              (c["exceso_pct"] or 0) / 100, (c["bajo_pct"] or 0) / 100] for c in n["colegios"]]
    fin = _tabla(ws, 4, ["Colegio", "Medidos", "Normal", "Sobrepeso", "Obesidad", "Bajo peso", "% exceso de peso", "% bajo peso"],
                 filas, [40, 11, 10, 12, 11, 11, 16, 12])
    for i in range(5, fin):
        ws.cell(i, 7).number_format = "0.0%"
        ws.cell(i, 8).number_format = "0.0%"
    ws.cell(fin + 1, 1, "Distribución").font = Font(bold=True, color=AZUL)
    for k_i, k in enumerate(CLASES):
        ws.cell(fin + 2 + k_i, 1, k)
        ws.cell(fin + 2 + k_i, 2, n["distribucion"][k])
    ch = DoughnutChart()
    ch.title, ch.height, ch.width = "Estado nutricional", 8, 12
    ch.add_data(Reference(ws, min_col=2, min_row=fin + 2, max_row=fin + 5))
    ch.set_categories(Reference(ws, min_col=1, min_row=fin + 2, max_row=fin + 5))
    ws.add_chart(ch, f"J4")

    ws2 = wb.create_sheet("IMC por edad")
    _titulo(ws2, "Promedios por edad y sexo", f"Año {n['anio']}", 8)
    _tabla(ws2, 4, ["Edad", "Estudiantes", "IMC niñas", "IMC niños", "Talla niñas (cm)", "Talla niños (cm)", "Peso niñas (kg)", "Peso niños (kg)"],
           [[e["edad"], e["n"], e["imc_F"], e["imc_M"], e["talla_F"], e["talla_M"], e["peso_F"], e["peso_M"]] for e in n["por_edad"]],
           [8, 12, 11, 11, 15, 15, 15, 15])

    ws3 = wb.create_sheet("Datos a revisar")
    _titulo(ws3, "Mediciones a revisar", "Valores que parecen mal escritos (no se cuentan en las estadísticas)", 6)
    _tabla(ws3, 4, ["Estudiante", "RUT", "Colegio", "Talla", "Peso", "IMC", "Problema"],
           [[r["nombre"], _rut(r["rut"]), r["colegio"], r["altura"], r["peso"], r["imc"], ", ".join(r["problemas"])] for r in n["a_revisar"]],
           [32, 14, 34, 9, 9, 9, 34])


def _hoja_deriv(wb, d):
    ws = wb.create_sheet("Derivaciones")
    _titulo(ws, "Derivaciones por especialidad",
            f"{d['derivados']} de {d['evaluados']} estudiantes evaluados tienen al menos una derivación", 3)
    filas = [[e["especialidad"], e["estudiantes"], (e["pct"] or 0) / 100] for e in d["especialidades"]]
    fin = _tabla(ws, 4, ["Especialidad", "Estudiantes derivados", "% de evaluados"], filas, [34, 20, 16])
    for i in range(5, fin):
        ws.cell(i, 3).number_format = "0.0%"
    if filas:
        ch = BarChart()
        ch.type, ch.title, ch.height, ch.width, ch.legend = "bar", "Estudiantes derivados", 9, 18, None
        ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=fin - 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fin - 1))
        ch.y_axis.majorGridlines = None
        ws.add_chart(ch, "E4")

    esps = [e["especialidad"] for e in d["especialidades"]]
    ws2 = wb.create_sheet("Derivaciones por colegio")
    _titulo(ws2, "Derivaciones por colegio y especialidad", "Cantidad de estudiantes derivados", len(esps) + 2)
    _tabla(ws2, 4, ["Colegio", "Derivados"] + esps,
           [[c["colegio"], c["derivados"]] + [c.get(e, 0) for e in esps] for c in d["colegios"]],
           [40, 11] + [14] * len(esps))

    ws3 = wb.create_sheet("Estudiantes derivados")
    _titulo(ws3, "Listado de estudiantes derivados", "Una fila por evaluación con derivación", 6)
    _tabla(ws3, 4, ["Estudiante", "RUT", "Colegio", "Evaluado por", "Fecha", "Derivado a"],
           [[x["nombre"], _rut(x["rut"]), x["colegio"], x["origen"], _f(x["fecha"]), ", ".join(x["especialidades"])] for x in d["lista"]],
           [32, 14, 34, 18, 12, 46])


def armar_excel(tipo, datos, filtros):
    wb = Workbook()
    _hoja_resumen(wb, datos, filtros)
    if "reev" in datos:
        _hoja_reev(wb, datos["reev"])
    if "nutri" in datos:
        _hoja_nutri(wb, datos["nutri"])
    if "deriv" in datos:
        _hoja_deriv(wb, datos["deriv"])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
